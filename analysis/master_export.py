"""Master analysis dataset: every record in stroke_analysis_master.csv with the original Codebook variables,
the derived analysis variables and all data added from the later extracts (diagnosis codes, encounters,
medications, Lab Tests 2/3/4/7/8). Written to outputs/stroke_master_dataset.xlsx (Arial). No MRN or DOB;
study_id is the only identifier.
"""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from .utils import OUT_DIR, PHI_COLUMNS, RESULTS, log

F = "Arial"
LAB_NAMES = {"Hb (g/dL)": "hb", "MCV (fL)": "mcv", "MCH (pg)": "mch", "MCHC (g/dL)": "mchc", "RDW-CV (%)": "rdw_cv",
             "RDW-SD (fL)": "rdw_sd", "Platelets (x10^9/L)": "platelets"}
DESCR = {
    "race4": "Race, 4 groups (White / Black / Asian / Other-unknown)",
    "smoking3": "Smoking: Never / Ever / Unknown",
    "migraine_any": "Any migraine diagnosis (codes 1, 2, 9) = 1 (PI decision)",
    "anemia_cat": "Anemia grade from Codebook Hb: None ≥12, Mild 10–11.9, Moderate 8–9.9, Severe <8 g/dL",
    "anemia_any": "Any anemia (Hb <12 g/dL)",
    "mcv_cat_lab": "MCV category from the dated lab MCV (within ±30 d of Hb)",
    "mcv_lab_dated": "MCV (fL) from Lab_Tests_3 nearest to the Hb date, within ±30 days (analysis MCV)",
    "mcv_lab_gap_days": "Days from Hb to that MCV (negative = before Hb)",
    "rdw_lab_dated": "RDW-CV (%) from Lab_Tests_7 nearest to the Hb date, within ±30 days",
    "rdw_lab_gap_days": "Days from Hb to that RDW-CV",
    "anaemia_type": "Anemia type by dated MCV (micro <80 / normo 80–100 / macro >100 fL)",
    "anaemia_pattern": "Laboratory anemia pattern: MCV + RDW-CV (>14.5% = high), Bessman classification",
    "anaemia_coded": "Anemia with / without a coded iron-deficiency diagnosis (±1 y of Hb)",
    "hb_timing": "Timing of Hb relative to stroke date",
    "y_isch": "Ischemic stroke vs no stroke (other stroke types missing)",
    "y_incident": "Stroke after the index date vs no stroke",
    "death_date": "Death date (Encounters extracts)",
    "last_contact_date": "Last encounter date (Encounters_8 Create Date)",
    "fu_start": "Follow-up start: later of index date and Hb date",
    "fu_end": "Follow-up end: stroke, death or last encounter",
    "fu_py": "Person-years of follow-up from fu_start",
    "fu_ev_any": "Stroke during follow-up",
    "fu_ev_isch": "Ischemic stroke during follow-up",
    "fu_died": "Died without stroke during follow-up",
    "death_1y_hb": "Died within 1 year of the Hb date",
    "serious_illness": "Cancer, heart failure, CKD, chronic liver disease or HIV recorded at any time",
}


DESCR.update({
    "y_type1": "Stroke subtype 1 (Codebook stroke_type = 1) vs no stroke", "y_type2": "Stroke subtype 2 vs no stroke",
    "y_type3": "Stroke subtype 3 vs no stroke", "y_type4": "Stroke subtype 4 vs no stroke",
    "y_type5": "Stroke subtype 5 vs no stroke", "y_type9": "Stroke subtype 9 (unknown) vs no stroke",
    "age45": "Age ≥45 years at index", "black": "Black race (race4)", "fib_any": "Fibroids present (any group)",
    "anticoag": "Anticoagulant use (antithrombotic = 2)", "ischaemic": "Stroke type ischemic (Codebook stroke_type = 1)",
    "hb_lab_nearest": "CBC hemoglobin (g/dL) from Lab_Tests_2 nearest to the Codebook Hb date (validation of hgb)",
})


def _lab_cols(df):
    out = {}
    for k, short in LAB_NAMES.items():
        if f"lab_{k}" in df:
            out[f"lab_{k}"] = f"{short}_lab_nearest"
            out[f"lab_{k} gap"] = f"{short}_lab_nearest_gap_days"
    return out


def build(df, d, t, raw_cols):
    from .followup import encounters
    m = df.copy()
    # keep the Codebook MCV under its original name; the dated lab MCV gets its own name
    if "mcv_codebook" in m:
        m["mcv_lab_dated"] = m["mcv"]
        m["mcv"] = m["mcv_codebook"]
        m = m.drop(columns=["mcv_codebook"])
    m = m.rename(columns={"mcv_gap": "mcv_lab_gap_days", "rdw": "rdw_lab_dated", "rdw_gap": "rdw_lab_gap_days",
                          "mcv_cat": "mcv_cat_lab"})
    m = m.rename(columns=_lab_cols(m))
    lc, dth = encounters()
    m["death_date"] = m.mrn.map(dth)
    m["last_contact_date"] = m.mrn.map(lc)
    # Paper 2 variables built on the eligible cohort (comorbidity flags, anemia type/pattern, outcomes)
    add = [c for c in d.columns if c not in m.columns and c not in df.columns and not c.startswith(("hb_below", "ht_"))
           and c not in ("hb_drop", "hb_above13", "mcv_codebook")]
    m = m.join(d[add])
    fu = t[["start", "end", "py", "ev_any", "ev_isch", "died"]].rename(columns=lambda c: f"fu_{c}")
    m = m.join(fu)
    # order: original Codebook columns, then everything added
    m = m.drop(columns=[c for c in m.columns if c.startswith("ht_") or c in ("age2",)])
    orig = [c for c in raw_cols if c in m.columns and c not in PHI_COLUMNS]
    rest = [c for c in m.columns if c not in orig and c not in PHI_COLUMNS]
    m = m[orig + rest]
    assert not [c for c in m.columns if c.lower() in PHI_COLUMNS]
    groups = {c: "Codebook (original)" for c in orig}
    for c in rest:
        groups[c] = ("Lab Tests extracts" if ("_lab_" in c or c.startswith(("mcv_lab", "rdw_lab"))) else
                     "Follow-up (Encounters)" if c.startswith("fu_") or c in ("death_date", "last_contact_date") else
                     "Paper 2 (diagnosis extracts / derived)" if c in add else "Derived (analysis)")
    return m, groups


def _cell(ws, v, font):
    c = WriteOnlyCell(ws, value=v)
    c.font = font
    return c


def write(df, d, t, raw_cols):
    m, groups = build(df, d, t, raw_cols)
    path = OUT_DIR / "stroke_master_dataset.xlsx"
    wb = Workbook(write_only=True)
    hdr_font, body = Font(name=F, bold=True, color="FFFFFF"), Font(name=F, size=10)
    fill = PatternFill("solid", fgColor="2A78D6")

    ws = wb.create_sheet("README")
    for row in [["Stroke & benign uterine disease — master analysis dataset"],
                [f"{len(m):,} records (all rows of stroke_analysis_master.csv), {m.shape[1]} variables. "
                 "study_id is the only identifier; MRN and DOB are not included."],
                ["Original Codebook variables are unchanged (including the Codebook mcv). Added variables are listed "
                 "in the 'Variables' sheet with their source."],
                ["Lab values: plausible value nearest to each woman's Hb date (blank, sentinel and implausible "
                 "results removed; see LAB_cleaning in stats_results.xlsx)."],
                ["Paper 2 and follow-up variables are filled for eligible women (eligible = 1) only."]]:
        ws.append([_cell(ws, row[0], body)])

    wv = wb.create_sheet("Variables")
    wv.append([_cell(wv, h, hdr_font) for h in ["Variable", "Source", "Description", "Non-missing"]])
    from .paper2_ext import LABELS
    for c in m.columns:
        wv.append([_cell(wv, c, body), _cell(wv, groups[c], body), _cell(wv, DESCR.get(c, LABELS.get(c, "")), body),
                   _cell(wv, int(m[c].notna().sum()), body)])

    wd = wb.create_sheet("Data")
    wd.freeze_panes = "B2"
    hdr = []
    for h in m.columns:
        c = WriteOnlyCell(wd, value=h)
        c.font, c.fill = hdr_font, fill
        hdr.append(c)
    wd.append(hdr)
    out = m.copy()
    for c in out.columns:
        if pd.api.types.is_datetime64_any_dtype(out[c]):
            out[c] = out[c].dt.strftime("%Y-%m-%d")
        elif out[c].dtype == object:
            out[c] = out[c].map(lambda v: v[:32000] if isinstance(v, str) else v)
    vals = out.astype(object).where(out.notna(), None).values.tolist()
    for r in vals:
        wd.append([_cell(wd, (v.item() if isinstance(v, np.generic) else v), body) for v in r])
    wb.save(path)
    RESULTS["master_export"] = dict(rows=len(m), cols=m.shape[1], path=path.name)
    log("Master", f"{path.name}: {len(m):,} records × {m.shape[1]} variables (no MRN/DOB)")
    return path
